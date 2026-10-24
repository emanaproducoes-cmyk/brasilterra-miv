"""Gera modelos de documento da Brasil Terra: assinatura de e-mail (HTML), .docx, .pptx e .xlsx."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Mm, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'public', 'downloads', '05_Papelaria_Digital')
LOGO = os.path.join(ROOT, 'public', 'downloads', '02_Logo')
L_H_FULL = os.path.join(LOGO, 'horizontal', 'fullcolor', 'BrasilTerra_horizontal_fullcolor_1024px.png')
L_H_NEG = os.path.join(LOGO, 'horizontal', 'negativa', 'BrasilTerra_horizontal_negativa_1024px.png')
L_S_FULL = os.path.join(LOGO, 'simbolo', 'fullcolor', 'BrasilTerra_simbolo_fullcolor_512px.png')
AZ, AM, PR, PROF, GR, GR2, BR, PAP, TER = '063E9B', 'FEBF26', '000000', '031B47', '262A33', '4B5260', 'E4E7EC', 'F6F4EE', 'B5532C'
rgb = lambda h: RGBColor.from_string(h)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor); tcPr.append(sh)

def cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for edge, (sz, color) in kw.items():
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:color'), color); b.append(e)
    tcPr.append(b)

def set_font(style, name, size, bold=False, color=GR, italic=False):
    style.font.name = name; style.font.size = Pt(size); style.font.bold = bold; style.font.italic = italic
    style.font.color.rgb = rgb(color)
    rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), name)

# ------------------------------------------------------------------ DOCX
def docx():
    d = Document()
    sec = d.sections[0]; sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20); sec.top_margin = Mm(28); sec.bottom_margin = Mm(22)
    st = d.styles
    set_font(st['Normal'], 'Source Sans 3', 10.5); st['Normal'].paragraph_format.space_after = Pt(7); st['Normal'].paragraph_format.line_spacing = 1.3
    set_font(st['Title'], 'Montserrat', 30, True, PR); st['Title'].paragraph_format.space_after = Pt(10)
    for lvl, size, color in [(1, 22, PR), (2, 15, AZ), (3, 12, AZ)]:
        s = st[f'Heading {lvl}']; set_font(s, 'Montserrat', size, True, color)
        s.paragraph_format.space_before = Pt(18 if lvl == 1 else 12); s.paragraph_format.space_after = Pt(6)
    s = st.add_style('BT Sobretítulo', WD_STYLE_TYPE.PARAGRAPH); s.base_style = st['Normal']; set_font(s, 'Montserrat', 8, True, AZ)
    s.font.all_caps = True; s.paragraph_format.space_after = Pt(2)
    s = st.add_style('BT Lide', WD_STYLE_TYPE.PARAGRAPH); s.base_style = st['Normal']; set_font(s, 'Source Sans 3', 14, False, GR2)
    s = st.add_style('BT Legenda', WD_STYLE_TYPE.PARAGRAPH); s.base_style = st['Normal']; set_font(s, 'Source Sans 3', 8, False, GR2)
    s = st.add_style('BT Número KPI', WD_STYLE_TYPE.PARAGRAPH); s.base_style = st['Normal']; set_font(s, 'Montserrat', 26, True, AZ)
    s.paragraph_format.space_after = Pt(0)
    s = st.add_style('BT Citação', WD_STYLE_TYPE.PARAGRAPH); s.base_style = st['Normal']; set_font(s, 'Source Sans 3', 14, False, AZ, italic=True)

    # cabeçalho / rodapé
    hp = sec.header.paragraphs[0]; hp.add_run().add_picture(L_H_FULL, width=Mm(50))
    fp = sec.footer.paragraphs[0]; fp.style = st['BT Legenda']
    fp.add_run('Brasil Terra · Consultoria e Projetos de Viabilidade · brasilterra.com.br    |    página ')
    r = fp.add_run(); f1 = OxmlElement('w:fldSimple'); f1.set(qn('w:instr'), 'PAGE'); r._r.append(f1)

    # capa
    d.add_paragraph('Relatório técnico · Projeto FNO', style='BT Sobretítulo')
    d.add_paragraph('Estudo de viabilidade técnico-financeira', style='Title')
    d.add_paragraph('[Nome do empreendimento] · [Município – UF]', style='BT Lide')
    d.add_paragraph('Versão para análise · [mês de ano]', style='BT Legenda')
    d.add_paragraph('Modelo de documento com os estilos nomeados da marca: Título, Título 1–3, BT Sobretítulo, BT Lide, BT Legenda, BT Número KPI, BT Citação. Substitua os textos entre colchetes.', style='BT Legenda')
    d.add_page_break()

    d.add_paragraph('01', style='BT Sobretítulo'); d.add_heading('Resumo executivo', 1)
    d.add_paragraph('Descreva em uma página o objetivo do investimento, o valor total, a fonte de recursos pretendida e a conclusão do estudo. '
                    'Use frases curtas e números com origem declarada.')
    t = d.add_table(rows=2, cols=3); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (n, l) in enumerate([('00,0%', 'TIR estimada'), ('0,0 anos', 'Payback'), ('0,00', 'ICSD médio')]):
        c = t.cell(0, i); shade(c, PAP); c.paragraphs[0].style = st['BT Número KPI']; c.paragraphs[0].text = n
        c2 = t.cell(1, i); shade(c2, PAP); c2.paragraphs[0].style = st['BT Legenda']; c2.paragraphs[0].text = l
    d.add_paragraph()
    d.add_heading('Projeção de resultados (R$ mil)', 2)
    rows = [('Item', 'Ano 1', 'Ano 2', 'Ano 3', 'Ano 4'), ('Receita bruta', '0,0', '0,0', '0,0', '0,0'),
            ('Custos variáveis', '(0,0)', '(0,0)', '(0,0)', '(0,0)'), ('Custos fixos', '(0,0)', '(0,0)', '(0,0)', '(0,0)'),
            ('Resultado operacional', '0,0', '0,0', '0,0', '0,0')]
    t = d.add_table(rows=len(rows), cols=5)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); p = c.paragraphs[0]; run = p.add_run(v); run.font.size = Pt(9)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            if i == 0: shade(c, PROF); run.font.color.rgb = rgb('FFFFFF'); run.font.bold = True; run.font.name = 'Montserrat'
            elif i == len(rows) - 1: shade(c, 'FFF2CC'); run.font.bold = True; cell_border(c, top=(12, PR))
            else:
                cell_border(c, bottom=(4, BR))
                if v.startswith('('): run.font.color.rgb = rgb(TER)
    d.add_paragraph('Tabela 1 · Valores ilustrativos. Negativos entre parênteses, em Terracota.', style='BT Legenda')

    d.add_heading('Parecer técnico', 2)
    t = d.add_table(rows=1, cols=1); c = t.cell(0, 0); shade(c, PAP); cell_border(c, left=(36, AM))
    c.paragraphs[0].add_run('Conclusão: ').bold = True
    c.paragraphs[0].add_run('[Projeto apto a ser submetido à análise da instituição financeira, condicionado a …]')
    d.add_paragraph('A análise e a aprovação do crédito são de responsabilidade exclusiva da instituição financeira.', style='BT Legenda')
    d.save(os.path.join(OUT, 'BrasilTerra_modelo_relatorio.docx'))

# ------------------------------------------------------------------ PPTX
def pptx():
    from pptx import Presentation
    from pptx.util import Inches, Pt as P, Emu
    from pptx.dml.color import RGBColor as C
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    pr = Presentation(); pr.slide_width, pr.slide_height = Inches(13.333), Inches(7.5)
    W, H = pr.slide_width, pr.slide_height; blank = pr.slide_layouts[6]
    def rect(s, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
        r = s.shapes.add_shape(shape, x, y, w, h); r.fill.solid(); r.fill.fore_color.rgb = C.from_string(color); r.line.fill.background(); return r
    def text(s, x, y, w, h, t, size, color, bold=False, font='Montserrat'):
        tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
        for i, line in enumerate(t.split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); r = p.add_run(); r.text = line
            r.font.size = P(size); r.font.bold = bold; r.font.name = font; r.font.color.rgb = C.from_string(color)
        return tb
    def bg(s, color): f = s.background.fill; f.solid(); f.fore_color.rgb = C.from_string(color)
    # 1 capa
    s = pr.slides.add_slide(blank); bg(s, PROF)
    for i, (x, y, w, h, c) in enumerate([(8.6, 0, 1.0, 4.2, '06286A'), (9.9, 0, 3.43, 2.0, AZ), (9.9, 2.2, 3.43, 2.0, '0A3380'), (8.6, 4.4, 4.73, 3.1, '06286A')]):
        rect(s, Inches(x), Inches(y), Inches(w), Inches(h), c)
    text(s, Inches(0.8), Inches(1.6), Inches(7.4), Inches(0.4), 'APRESENTAÇÃO INSTITUCIONAL', 12, AM, True)
    text(s, Inches(0.8), Inches(2.1), Inches(7.4), Inches(2.6), 'Projetos que transformam\nambição em crédito analisável.', 36, 'FFFFFF', True)
    text(s, Inches(0.8), Inches(4.6), Inches(7), Inches(0.4), '[Cliente] · [mês de ano]', 14, 'AAB4C8', False, 'Source Sans 3')
    s.shapes.add_picture(L_H_NEG, Inches(0.8), Inches(6.2), width=Inches(2.6))
    # 2 divisor
    s = pr.slides.add_slide(blank); bg(s, PROF)
    text(s, Inches(0.8), Inches(1.4), Inches(6), Inches(0.5), 'DIAGNÓSTICO', 16, AM, True)
    text(s, Inches(0.8), Inches(1.9), Inches(8), Inches(3.4), '01.', 160, AM, True)
    rect(s, 0, H - Inches(0.14), Inches(5), Inches(0.14), AM)
    # 3 conteúdo
    s = pr.slides.add_slide(blank); bg(s, PAP)
    rect(s, Inches(0.8), Inches(0.9), Inches(0.6), Inches(0.06), AM)
    text(s, Inches(0.8), Inches(1.0), Inches(8), Inches(0.4), '01 · DIAGNÓSTICO', 11, AZ, True)
    text(s, Inches(0.8), Inches(1.4), Inches(11), Inches(1.0), 'Título do slide em uma linha', 32, PR, True)
    text(s, Inches(0.8), Inches(2.6), Inches(7.2), Inches(3.5), '• Tópico curto, com no máximo duas linhas\n• Números sempre com fonte e unidade\n• Um destaque Amarelo Safra por slide', 18, GR, False, 'Source Sans 3')
    s.shapes.add_picture(L_S_FULL, W - Inches(1.3), H - Inches(1.1), width=Inches(0.5))
    # 4 KPIs
    s = pr.slides.add_slide(blank); bg(s, PAP)
    text(s, Inches(0.8), Inches(1.0), Inches(8), Inches(0.4), '03 · RESUMO EXECUTIVO', 11, AZ, True)
    text(s, Inches(0.8), Inches(1.4), Inches(11), Inches(1.0), 'O projeto em números', 32, PR, True)
    for i, (n, l) in enumerate([('00,0%', 'TIR estimada*'), ('0,0', 'anos de payback*'), ('0,00', 'ICSD médio*'), ('00', 'empregos diretos*')]):
        x = Inches(0.8 + i * 3.0); rect(s, x, Inches(3.0), Inches(2.8), Inches(2.2), PROF)
        text(s, x + Inches(0.3), Inches(3.3), Inches(2.4), Inches(1.0), n, 44, AM, True)
        text(s, x + Inches(0.3), Inches(4.4), Inches(2.4), Inches(0.5), l, 14, 'FFFFFF', False, 'Source Sans 3')
    text(s, Inches(0.8), Inches(6.4), Inches(11), Inches(0.4), '*Valores ilustrativos. A aprovação do crédito é de responsabilidade exclusiva da instituição financeira.', 11, GR2, False, 'Source Sans 3')
    # 5 encerramento
    s = pr.slides.add_slide(blank); bg(s, PROF)
    s.shapes.add_picture(L_H_NEG, Inches(4.17), Inches(2.6), width=Inches(5))
    text(s, Inches(0.8), Inches(5.6), Inches(11.7), Inches(0.5), 'brasilterra.com.br · [telefone] · [e-mail]', 14, 'C8D0DE', False, 'Source Sans 3')
    pr.save(os.path.join(OUT, 'BrasilTerra_modelo_apresentacao.pptx'))

# ------------------------------------------------------------------ XLSX
def xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
    from openpyxl.drawing.image import Image as XImg
    wb = Workbook(); ws = wb.active; ws.title = 'Fluxo de caixa'
    head = NamedStyle('BT Cabeçalho'); head.font = Font(name='Montserrat', bold=True, color='FFFFFF', size=10)
    head.fill = PatternFill('solid', fgColor=PROF); head.alignment = Alignment(horizontal='center', vertical='center')
    num = NamedStyle('BT Valor'); num.font = Font(name='Source Sans 3', size=10, color=GR); num.number_format = '#,##0.00;[Color53](#,##0.00)'
    num.border = Border(bottom=Side(style='thin', color=BR))
    lab = NamedStyle('BT Rótulo'); lab.font = Font(name='Source Sans 3', size=10, color=GR); lab.border = Border(bottom=Side(style='thin', color=BR))
    tot = NamedStyle('BT Total'); tot.font = Font(name='Source Sans 3', bold=True, size=10, color=PR); tot.fill = PatternFill('solid', fgColor='FFF2CC')
    tot.number_format = '#,##0.00;(#,##0.00)'; tot.border = Border(top=Side(style='medium', color=PR))
    for s in (head, num, lab, tot): wb.add_named_style(s)
    ws.row_dimensions[1].height = 60
    ws.add_image(XImg(L_H_FULL), 'A1'); ws['A1'].value = None
    ws._images[0].width, ws._images[0].height = 200, 59
    ws['A4'] = 'Projeção de fluxo de caixa (R$) · modelo'; ws['A4'].font = Font(name='Montserrat', bold=True, size=14, color=PR)
    ws['A5'] = 'Preencha as células em azul-claro. Totais e indicadores são calculados automaticamente.'; ws['A5'].font = Font(name='Source Sans 3', size=9, color=GR2)
    cols = ['Item'] + [f'Ano {i}' for i in range(0, 9)]
    for j, c in enumerate(cols): ws.cell(7, j + 1, c).style = 'BT Cabeçalho'
    items = ['Investimento fixo', 'Capital de giro', 'Receita bruta', 'Custos variáveis', 'Custos fixos', 'Impostos']
    inp = PatternFill('solid', fgColor='E8F1FB')
    for i, it in enumerate(items):
        r = 8 + i; ws.cell(r, 1, it).style = 'BT Rótulo'
        for j in range(2, 11): c = ws.cell(r, j, 0); c.style = 'BT Valor'; c.fill = inp
    r = 8 + len(items); ws.cell(r, 1, 'Fluxo líquido').style = 'BT Total'
    for j in range(2, 11):
        L = ws.cell(8, j).column_letter
        ws.cell(r, j, f'=-{L}8-{L}9+{L}10-{L}11-{L}12-{L}13').style = 'BT Total'
    ws.cell(r + 2, 1, 'Taxa mínima de atratividade (a.a.)').style = 'BT Rótulo'; c = ws.cell(r + 2, 2, 0.12); c.number_format = '0.0%'; c.fill = inp
    ws.cell(r + 3, 1, 'VPL').style = 'BT Rótulo'; c = ws.cell(r + 3, 2, f'=B{r}+NPV(B{r+2},C{r}:J{r})'); c.style = 'BT Valor'
    ws.cell(r + 4, 1, 'TIR').style = 'BT Rótulo'; c = ws.cell(r + 4, 2, f'=IFERROR(IRR(B{r}:J{r}),"—")'); c.number_format = '0.0%'
    ws.cell(r + 6, 1, 'Valores ilustrativos. A aprovação do crédito é de responsabilidade exclusiva da instituição financeira.').font = Font(name='Source Sans 3', size=8, italic=True, color=GR2)
    ws.column_dimensions['A'].width = 34
    for j in range(2, 11): ws.column_dimensions[ws.cell(7, j).column_letter].width = 14
    ws.freeze_panes = 'B8'; ws.sheet_view.showGridLines = False
    wb.save(os.path.join(OUT, 'BrasilTerra_modelo_fluxo_de_caixa.xlsx'))

# ------------------------------------------------------------------ Assinatura
SIG = '''<!DOCTYPE html>
<html lang="pt-BR"><head><meta charset="utf-8"><title>Assinatura de e-mail · Brasil Terra</title></head>
<body style="margin:0;padding:24px;background:#ffffff;">
<!--
  Assinatura de e-mail Brasil Terra · compatível com Outlook, Gmail e Apple Mail.
  1. Hospede o arquivo brasilterra-simbolo-128.png (pasta ao lado) no site da empresa e troque o endereço em src="...".
  2. Substitua os campos entre colchetes.
  3. Abra este arquivo no navegador, selecione tudo (Ctrl+A), copie e cole no campo de assinatura do cliente de e-mail.
-->
<table role="presentation" cellpadding="0" cellspacing="0" border="0" width="520" style="width:520px;max-width:100%;border-collapse:collapse;font-family:Arial,Helvetica,sans-serif;">
  <tr><td colspan="2" style="border-top:3px solid #FEBF26;font-size:0;line-height:0;height:14px;">&nbsp;</td></tr>
  <tr>
    <td valign="top" width="76" style="width:76px;padding:0 16px 0 0;">
      <img src="https://brasilterra.com.br/assinatura/brasilterra-simbolo-128.png" width="64" height="60" alt="Brasil Terra" style="display:block;border:0;width:64px;height:60px;">
    </td>
    <td valign="top" style="font-family:Arial,Helvetica,sans-serif;color:#262A33;font-size:13px;line-height:19px;">
      <strong style="font-size:15px;color:#000000;">[Nome Sobrenome]</strong><br>
      <span style="color:#4B5260;">[Cargo]</span><br>
      <span style="font-size:11px;font-weight:bold;letter-spacing:0.5px;color:#063E9B;">BRASIL TERRA · Consultoria e Projetos de Viabilidade</span><br>
      <a href="tel:+5591000000000" style="color:#262A33;text-decoration:none;">[+55 91 00000-0000]</a> ·
      <a href="https://wa.me/5591000000000" style="color:#063E9B;text-decoration:none;">WhatsApp</a><br>
      <a href="mailto:[nome]@brasilterra.com.br" style="color:#262A33;text-decoration:none;">[nome]@brasilterra.com.br</a> ·
      <a href="https://brasilterra.com.br" style="color:#063E9B;text-decoration:none;">brasilterra.com.br</a><br>
      <span style="color:#4B5260;">[Cidade] · [UF]</span>
    </td>
  </tr>
  <tr><td colspan="2" style="padding-top:12px;font-family:Arial,Helvetica,sans-serif;font-size:10px;line-height:14px;color:#6B7280;">
    Esta mensagem pode conter informação confidencial. Se você a recebeu por engano, avise o remetente e apague-a.
    A Brasil Terra estrutura projetos para análise; a aprovação do crédito é de responsabilidade exclusiva da instituição financeira.
  </td></tr>
</table>
</body></html>
'''

if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    open(os.path.join(OUT, 'assinatura_email.html'), 'w', encoding='utf-8').write(SIG)
    from PIL import Image
    im = Image.open(L_S_FULL); im.thumbnail((128, 128)); im.save(os.path.join(OUT, 'brasilterra-simbolo-128.png'))
    docx(); pptx(); xlsx(); print('modelos ok')
